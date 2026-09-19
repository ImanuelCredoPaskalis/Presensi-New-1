"""
Module: export_manager.py
Ekspor data laporan presensi ke Excel (.xlsx) dengan openpyxl dan format CSV.
"""

import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_to_excel(df: pd.DataFrame, title: str = "Laporan Presensi Mahasiswa Lab FisMat") -> io.BytesIO:
    """Menghasilkan file Excel (.xlsx) dengan format dan styling rapi."""
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap_Presensi"
    ws.views.sheetView[0].showGridLines = True

    # Style definitions
    font_title = Font(name="Arial", size=14, bold=True, color="003640")
    font_subtitle = Font(name="Arial", size=10, italic=True, color="555555")
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Arial", size=9)
    
    fill_header = PatternFill(start_color="171B26", end_color="171B26", fill_type="solid")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    # 1. Judul Dokumen
    ws.append([title.upper()])
    ws.cell(row=1, column=1).font = font_title
    ws.append(["Laboratorium Micro Teaching Fisika & Matematika • Sistem Presensi Web Edition"])
    ws.cell(row=2, column=1).font = font_subtitle
    ws.append([]) # Baris kosong

    # 2. Header Tabel
    headers = list(df.columns)
    ws.append(headers)
    header_row_idx = 4
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[header_row_idx].height = 24

    # 3. Baris Data
    for r_idx, row in enumerate(df.itertuples(index=False), start=header_row_idx + 1):
        ws.append(list(row))
        ws.row_dimensions[r_idx].height = 20
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = font_data
            cell.border = thin_border
            # Zebra striping
            if r_idx % 2 == 0:
                cell.fill = fill_alt
            # Align center untuk kolom waktu & status
            val_str = str(cell.value)
            if any(k in headers[c_idx-1].lower() for k in ["jam", "tanggal", "status", "durasi", "no"]):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # 4. Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row >= header_row_idx and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output)
    output.seek(0)
    return output


def export_to_csv(df: pd.DataFrame) -> bytes:
    """Mengembalikan data dalam bentuk CSV bytes."""
    return df.to_csv(index=False).encode("utf-8")
